from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile
import unittest

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

from ppt_automator.embedded_workbook_writer import update_embedded_workbook


class EmbeddedWorkbookWriterTypedTests(unittest.TestCase):
    def test_typed_matrix_writes_text_cells_numbers_and_resizes_table(self) -> None:
        original = _workbook_bytes()
        matrix = [
            [
                {"value": "", "type": "text", "force_text": True, "source_raw": ""},
                {"value": "TIM", "type": "text", "force_text": True, "source_raw": ""},
                {"value": "001", "type": "text", "force_text": True, "source_raw": ""},
            ],
            [
                {"value": "Nov/25", "type": "text", "force_text": True, "source_raw": ""},
                {"value": "45.8419", "type": "number", "force_text": False, "source_raw": "45.8419"},
                {"value": "43", "type": "number", "force_text": False, "source_raw": "43"},
            ],
        ]

        updated = update_embedded_workbook(original, "Planilha1", matrix)

        workbook = openpyxl.load_workbook(BytesIO(updated), data_only=True)
        worksheet = workbook["Planilha1"]
        self.assertEqual(worksheet["A2"].value, "Nov/25")
        self.assertEqual(worksheet["C1"].value, "001")
        self.assertAlmostEqual(worksheet["B2"].value, 45.8419)
        self.assertEqual(worksheet.tables["Tabela1"].ref, "A1:C2")
        workbook.close()

    def test_preserves_package_parts_and_changes_only_expected_xlsx_parts(self) -> None:
        original = _workbook_bytes()
        updated = update_embedded_workbook(original, "Planilha1", [[" ", "TOTAL"], ["Cristal", 15.990453460620525]])

        with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(updated)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
            self.assertLessEqual(set(changed), {"xl/worksheets/sheet1.xml", "xl/tables/table1.xml", "xl/sharedStrings.xml"})


def _workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Planilha1"
    for row in [[" ", "Old"], ["A", 1]]:
        worksheet.append(row)
    table = Table(displayName="Tabela1", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    worksheet.add_table(table)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


if __name__ == "__main__":
    unittest.main()
